"""엑셀 리포트 출력 (openpyxl).

pandas 를 쓰지 않는다 — 버전에 따라 to_excel 동작이 달라지고,
셀 단위 서식·조건부서식을 직접 다루는 편이 결과물이 훨씬 읽기 좋다.

시트 구성
  설명          산식 정의, 실행 파라미터, 환율, 주의사항  ← 먼저 읽어야 하는 시트
  TOP30_통합     5개 지역 ROIC 상위 30을 한 장에
  <지역>_TOP30   지역별 상위 30
  계산근거       NOPAT·투하자본·실효세율·사용 보고기간·경고
  시총상위100    유니버스 전체 (제외 여부 포함)
  제외종목       제외 사유별
  조회실패       티커 오류·상장폐지·수집 실패
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import REGIONS
from .screener import Row, ScreenResult

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13, color="1F3864")
SUB_FONT = Font(bold=True, size=11, color="1F3864")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INT_CUR = {"KRW", "JPY"}  # 소수점 없이 표시할 통화


# ---------------------------------------------------------------------------
# 메인 랭킹 시트 컬럼 정의
# ---------------------------------------------------------------------------
def _price_fmt(currency: str | None) -> str:
    return "#,##0" if (currency or "").upper() in INT_CUR else "#,##0.00"


COLUMNS: list[tuple[str, int, str]] = [
    # (헤더, 폭, 숫자서식)  서식이 'price' 면 통화별로 결정
    ("ROIC순위", 8, "0"),
    ("지역", 6, ""),
    ("티커", 12, ""),
    ("기업명", 30, ""),
    ("업종(섹터)", 20, ""),
    ("세부업종", 26, ""),
    ("ROIC", 9, "0.0%"),
    ("EV/EBITDA", 11, "0.0"),
    ("현재가", 12, "price"),
    ("52주최고", 12, "price"),
    ("52주최저", 12, "price"),
    ("역대최고", 12, "price"),
    ("역대최고일", 12, "yyyy-mm-dd"),
    ("역대최저", 12, "price"),
    ("역대최저일", 12, "yyyy-mm-dd"),
    ("현재가/역대최고", 14, "0.0%"),
    ("지역시총순위", 12, "0"),
    ("시가총액(현지)", 18, "#,##0"),
    ("시가총액(USD,백만)", 18, "#,##0"),
    ("통화", 7, ""),
    ("거래소", 20, ""),
    ("매매단위(참고)", 14, ""),
    ("미래에셋온라인(참고)", 20, ""),
    ("기간커버", 22, ""),
    ("주의", 40, ""),
]


def _row_values(r: Row, rank: int) -> list:
    res = r.roic
    p = r.price
    region = REGIONS[r.region]
    coverage = ""
    if res and res.coverage_start and res.coverage_end:
        coverage = (
            f"{res.coverage_start.isoformat()}~{res.coverage_end.isoformat()}"
            f" ({res.coverage_ratio:.0%})"
        )
    return [
        rank,
        region.name,
        r.ticker,
        r.name,
        r.sector,
        r.industry,
        r.roic_value,
        r.ev_to_ebitda,
        p.current,
        p.high_52w,
        p.low_52w,
        p.high_all,
        p.high_all_date,
        p.low_all,
        p.low_all_date,
        p.pct_of_all_high,
        r.mcap_rank,
        r.market_cap,
        (r.market_cap_usd / 1e6) if r.market_cap_usd else None,
        r.currency,
        r.exchange,
        region.typical_lot,
        region.mirae_online,
        coverage,
        "; ".join(res.warnings) if res and res.warnings else "",
    ]


def _write_table(
    ws: Worksheet,
    columns: list[tuple[str, int, str]],
    rows: list[list],
    *,
    freeze: str = "D2",
) -> None:
    for c, (header, width, _fmt) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30

    currency_idx = next((i for i, c in enumerate(columns) if c[0] == "통화"), None)
    for r_i, values in enumerate(rows, 2):
        currency = values[currency_idx] if currency_idx is not None else None
        for c_i, (value, (_h, _w, fmt)) in enumerate(zip(values, columns), 1):
            cell = ws.cell(row=r_i, column=c_i, value=value)
            cell.border = BORDER
            if fmt == "price":
                cell.number_format = _price_fmt(currency)
            elif fmt:
                cell.number_format = fmt
            if isinstance(value, str) and len(value) > 40:
                cell.alignment = Alignment(wrap_text=False)

    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws.freeze_panes = freeze


def _colorscale_on(ws: Worksheet, header: str, n_rows: int, columns=COLUMNS) -> None:
    idx = next((i for i, c in enumerate(columns, 1) if c[0] == header), None)
    if not idx or n_rows < 2:
        return
    col = get_column_letter(idx)
    ws.conditional_formatting.add(
        f"{col}2:{col}{n_rows + 1}",
        ColorScaleRule(
            start_type="min", start_color="FFF2CC",
            mid_type="percentile", mid_value=50, mid_color="9BC2E6",
            end_type="max", end_color="2E75B6",
        ),
    )


# ---------------------------------------------------------------------------
def _sheet_explain(wb: Workbook, result: ScreenResult) -> None:
    ws = wb.create_sheet("설명")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    p = result.params
    rp = p.roic

    cursor = {"row": 1}

    def line(label: str, value: str = "", *, style: Font | None = None) -> None:
        cursor["row"] += 1
        row = cursor["row"]
        a = ws.cell(row=row, column=1, value=label)
        a.font = style or Font(bold=True, size=10)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=row, column=2, value=value)
        b.alignment = Alignment(vertical="top", wrap_text=True)

    ws["A1"] = "ROIC 상위기업 스크리너 결과"
    ws["A1"].font = TITLE_FONT
    line("생성시각", result.generated_at.strftime("%Y-%m-%d %H:%M:%S"))
    line("")
    line("■ ROIC 정의", "", style=SUB_FONT)
    line("측정기간", f"{p.window_start} ~ {p.window_end} 구간을 하나의 창으로 보고 누적 계산")
    line(
        "산식",
        "ROIC = (Σ EBIT × (1 − 실효세율) ÷ 커버연수) ÷ 평균 투하자본",
    )
    line(
        "실효세율",
        f"Σ법인세비용 ÷ Σ세전이익, {rp.tax_rate_min:.0%}~{rp.tax_rate_max:.0%} 로 클램프. "
        "계산 불가 시 지역 법정세율 대체",
    )
    line(
        "투하자본",
        "총차입금 + 자본총계(비지배지분 포함) − 초과현금"
        + (" − 영업권" if rp.exclude_goodwill else "")
        + f". 초과현금 = max(0, 현금및단기금융상품 − 매출×{rp.operating_cash_pct:.0%})",
    )
    line(
        "평균 투하자본",
        "창 구간 전후의 모든 재무상태표 스냅샷 평균(직전 기말 포함). "
        "기말 한 시점만 쓰면 대형 인수·자사주매입 직후 ROIC가 왜곡된다",
    )
    line(
        "회계연도 정렬",
        "일본처럼 3월 결산이면 실제 커버 구간이 창과 어긋난다. "
        "각 종목의 '기간커버' 컬럼에 실측 구간과 커버율을 그대로 표시했다",
    )
    line("")
    line("■ 스크리닝 조건", "", style=SUB_FONT)
    line("1단계", f"지역별 시가총액(USD 환산) 상위 {p.top_n_marketcap}종목")
    line(
        "2단계",
        "금융·부동산(리츠)·유틸리티 제외 — 투하자본 개념이 성립하지 않아 ROIC 비교가 무의미",
    )
    line(
        "3단계",
        f"평균 투하자본 ${rp.min_invested_capital_usd / 1e6:.0f}M 미만 제외, "
        f"기간 커버리지 {rp.min_coverage_ratio:.0%} 미만 제외",
    )
    line("4단계", f"ROIC 상위 {p.top_n_roic}종목 선정")
    line("")
    line("■ 데이터 출처와 한계", "", style=SUB_FONT)
    line("재무·주가", "Yahoo Finance (yfinance). 비공식 래퍼이므로 계정 라벨 변경·IP 차단 가능")
    line(
        "유니버스",
        "무료로 '전 종목 시총 순위'를 주는 API 는 중국(akshare)뿐이다. "
        "나머지 4개 지역은 seeds/*.txt 시드 리스트 안에서의 시총 상위 100이다 — "
        "시드에 없는 대형주는 후보에 들어오지 못한다",
    )
    line(
        "역대 최고/최저가",
        "액면분할만 보정한 종가 기준(배당 보정 안 함). "
        "Yahoo 히스토리 시작일 이전 데이터는 반영되지 않으며, 증자·감자 보정 누락 가능",
    )
    line(
        "EV/EBITDA",
        "Yahoo enterpriseToEbitda 우선, 없으면 (시총 + 순차입금) ÷ 연율화 EBITDA 직접계산. "
        "'계산근거' 시트에 종목별 산출 방식 표기",
    )
    line("")
    line("■ 환율 (USD 환산 계수, 1단위당 USD)", "", style=SUB_FONT)
    for cur, rate in sorted(result.fx.items()):
        line(cur, f"{rate:.6f}" if rate else "조회 실패")
    if result.notes:
        line("")
        line("■ 실행 메모", "", style=SUB_FONT)
        for n in result.notes:
            line("", n)
    line("")
    line("■ 면책", "", style=SUB_FONT)
    line(
        "",
        "투자자문이 아니다. 무료 소스의 자동 수집 결과이며 검증되지 않은 값이 섞일 수 있다. "
        "매수 판단 전 원본 공시(DART/SEC/EDINET/ESEF)로 교차확인할 것.",
    )


def _sheet_ranking(wb: Workbook, title: str, rows: list[Row]) -> None:
    ws = wb.create_sheet(title)
    data = [_row_values(r, i) for i, r in enumerate(rows, 1)]
    _write_table(ws, COLUMNS, data)
    _colorscale_on(ws, "ROIC", len(data))


EVIDENCE_COLUMNS: list[tuple[str, int, str]] = [
    ("지역", 6, ""),
    ("티커", 12, ""),
    ("기업명", 28, ""),
    ("ROIC", 9, "0.0%"),
    ("연율화 NOPAT(현지)", 20, "#,##0"),
    ("누적 EBIT(현지)", 18, "#,##0"),
    ("누적 매출(현지)", 18, "#,##0"),
    ("누적 EBITDA(현지)", 18, "#,##0"),
    ("실효세율", 10, "0.0%"),
    ("세율근거", 20, ""),
    ("평균 투하자본(현지)", 20, "#,##0"),
    ("IC 스냅샷수", 12, "0"),
    ("총차입금(기말)", 18, "#,##0"),
    ("자본총계(기말)", 18, "#,##0"),
    ("초과현금(기말)", 18, "#,##0"),
    ("영업권(기말)", 18, "#,##0"),
    ("커버 시작", 12, ""),
    ("커버 종료", 12, ""),
    ("커버일수", 10, "0"),
    ("커버율", 9, "0%"),
    ("사용 보고기간", 46, ""),
    ("EV/EBITDA 산출", 26, ""),
    ("경고", 46, ""),
    ("통화", 7, ""),
]


def _sheet_evidence(wb: Workbook, rows: list[Row]) -> None:
    ws = wb.create_sheet("계산근거")
    data = []
    for r in rows:
        res = r.roic
        if res is None:
            continue
        data.append([
            REGIONS[r.region].name,
            r.ticker,
            r.name,
            res.roic,
            res.nopat_annualized,
            res.ebit_total,
            res.revenue_total,
            res.ebitda_total,
            res.tax_rate,
            res.tax_rate_source,
            res.invested_capital_avg,
            res.ic_snapshots,
            res.total_debt_last,
            res.equity_last,
            res.excess_cash_last,
            res.goodwill_last,
            res.coverage_start.isoformat() if res.coverage_start else None,
            res.coverage_end.isoformat() if res.coverage_end else None,
            res.coverage_days,
            res.coverage_ratio,
            ", ".join(res.periods_used),
            r.ev_to_ebitda_source,
            "; ".join(res.warnings),
            r.currency,
        ])
    _write_table(ws, EVIDENCE_COLUMNS, data, freeze="C2")


UNIVERSE_COLUMNS: list[tuple[str, int, str]] = [
    ("지역시총순위", 12, "0"),
    ("지역", 6, ""),
    ("티커", 12, ""),
    ("기업명", 30, ""),
    ("업종(섹터)", 20, ""),
    ("세부업종", 26, ""),
    ("시가총액(USD,백만)", 18, "#,##0"),
    ("시가총액(현지)", 18, "#,##0"),
    ("통화", 7, ""),
    ("거래소", 20, ""),
    ("ROIC", 9, "0.0%"),
    ("EV/EBITDA", 11, "0.0"),
    ("상태", 46, ""),
]


def _sheet_universe(wb: Workbook, result: ScreenResult) -> None:
    ws = wb.create_sheet(f"시총상위{result.params.top_n_marketcap}")
    selected = {(r.region, r.ticker) for rows in result.top.values() for r in rows}
    data = []
    for region in result.universe:
        for r in result.universe[region]:
            if (r.region, r.ticker) in selected:
                status = "★ ROIC 상위 선정"
            elif r.excluded_reason:
                status = r.excluded_reason
            else:
                status = "유니버스 통과 / ROIC 상위 미달"
            data.append([
                r.mcap_rank,
                REGIONS[r.region].name,
                r.ticker,
                r.name,
                r.sector,
                r.industry,
                (r.market_cap_usd / 1e6) if r.market_cap_usd else None,
                r.market_cap,
                r.currency,
                r.exchange,
                r.roic_value,
                r.ev_to_ebitda,
                status,
            ])
    _write_table(ws, UNIVERSE_COLUMNS, data, freeze="D2")


EXCLUDED_COLUMNS: list[tuple[str, int, str]] = [
    ("지역", 6, ""),
    ("티커", 12, ""),
    ("기업명", 30, ""),
    ("업종(섹터)", 20, ""),
    ("세부업종", 26, ""),
    ("지역시총순위", 12, "0"),
    ("시가총액(USD,백만)", 18, "#,##0"),
    ("제외사유", 70, ""),
]


def _sheet_excluded(wb: Workbook, result: ScreenResult) -> None:
    ws = wb.create_sheet("제외종목")
    data = [
        [
            REGIONS[r.region].name,
            r.ticker,
            r.name,
            r.sector,
            r.industry,
            r.mcap_rank,
            (r.market_cap_usd / 1e6) if r.market_cap_usd else None,
            r.excluded_reason,
        ]
        for r in result.excluded
    ]
    _write_table(ws, EXCLUDED_COLUMNS, data, freeze="C2")


def _sheet_errors(wb: Workbook, result: ScreenResult) -> None:
    ws = wb.create_sheet("조회실패")
    cols = [("티커", 14, ""), ("단계", 12, ""), ("메시지", 110, "")]
    data = [[e.ticker, e.stage, e.message] for e in result.errors]
    _write_table(ws, cols, data, freeze="A2")
    if data:
        for r_i in range(2, len(data) + 2):
            ws.cell(row=r_i, column=1).fill = WARN_FILL


# ---------------------------------------------------------------------------
def write_report(result: ScreenResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    _sheet_explain(wb, result)

    combined: list[Row] = []
    for region in result.top:
        combined.extend(result.top[region])
    # 통합 시트는 지역 무관 전체 ROIC 순위로 정렬한다(지역별 순위는 지역 시트에서 본다)
    combined.sort(key=lambda r: (r.roic_value is None, -(r.roic_value or 0.0)))
    if combined:
        _sheet_ranking(wb, f"TOP{result.params.top_n_roic}_통합", combined)
    for region, rows in result.top.items():
        _sheet_ranking(wb, f"{region.upper()}_TOP{result.params.top_n_roic}", rows)

    _sheet_evidence(wb, combined)
    _sheet_universe(wb, result)
    _sheet_excluded(wb, result)
    _sheet_errors(wb, result)

    wb.save(path)
    return path


def default_filename(regions: list[str], now: dt.datetime | None = None) -> str:
    now = now or dt.datetime.now()
    return f"roic_top30_{'-'.join(regions)}_{now:%Y%m%d_%H%M}.xlsx"
